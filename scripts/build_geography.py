"""Rebuild a GeoNames snapshot from explicit downloads and training catalog only.

Does not inspect news, labels, validation, test or LiveDay. Output defaults to a
private proposal, never silently overwrites the frozen production snapshot.
"""
from pathlib import Path
import json,zipfile,unicodedata,re,hashlib
from datetime import datetime,timezone
import argparse
from openpyxl import load_workbook
parser=argparse.ArgumentParser()
parser.add_argument('--geonames-dir',type=Path,default=Path('.runtime/geonames'))
parser.add_argument('--output',type=Path,default=Path('reports/private/geography-rebuild.json'))
args=parser.parse_args()
p=args.geonames_dir
def norm(s):return re.sub(r'[^a-z0-9]','',unicodedata.normalize('NFKD',s).encode('ascii','ignore').decode().lower())
countries={}
for line in (p/'countryInfo.txt').read_text().splitlines():
 if line.startswith('#'):continue
 v=line.split('\t')
 if len(v)>16:
  countries[v[0]]={'code':v[0],'name':v[4],'geoname_id':v[16],'iso3':v[1]}
admin={}
for line in (p/'admin1CodesASCII.txt').read_text().splitlines():
 v=line.split('\t');admin[v[0]]={'name':v[1],'ascii_name':v[2],'geoname_id':v[3]}
with zipfile.ZipFile(p/'cities15000.zip') as z:
 cityrows=[line.split('\t') for line in z.read('cities15000.txt').decode().splitlines()]
workbook=load_workbook('inputs/supplier_training.xlsx',read_only=True,data_only=True)
rows=workbook['Suppliers'].iter_rows(values_only=True);headers=next(rows)
columns={k:headers.index(k) for k in ('Country','City')}
locations=[{k:r[i] for k,i in columns.items()} for r in rows if r[0]]
workbook.close()
out={}
for loc in locations:
 country=loc['Country'].split(' (')[0]
 match=[c for c in countries.values() if norm(country) in {norm(c['name']),norm(c['name'].removeprefix('The ')),norm(c['iso3']),norm(c['code'])}]
 assert len(match)==1,(country,len(match))
 country=match[0];key=loc['Country']+'|'+loc['City']
 if key in out:continue
 matches=[]
 for v in cityrows:
  if v[8]!=country['code']:continue
  names=[v[1],v[2]]+v[3].split(',')
  if norm(loc['City']) not in {norm(n) for n in names}:continue
  adm=admin.get(v[8]+'.'+v[10])
  # Gazetteer alternatives describe places, never extra company facilities.
  aliases=sorted({n for n in names if len(n)>=4 and len(n)<65 and not n.startswith('http')})
  matches.append({'geoname_id':v[0],'name':v[1],'aliases':aliases,'admin1':adm,'source_url':'https://www.geonames.org/'+v[0]+'/','modified_at':v[18]})
 out[key]={'country':country,'input_city':loc['City'],'status':'ambiguous' if len(matches)>1 else 'matched' if matches else 'unresolved','places':matches}
 print(key, len(matches),[(x['name'],x['admin1']['name'] if x['admin1'] else None) for x in matches])
sources=[{'url':'https://download.geonames.org/export/dump/'+f,'sha256':hashlib.sha256((p/f).read_bytes()).hexdigest()} for f in ['cities15000.zip','admin1CodesASCII.txt','countryInfo.txt']]
data={'version':'geo-1.0.0','retrieved_at':datetime.now(timezone.utc).isoformat(),'attribution':'GeoNames, CC BY 4.0','documentation':'https://download.geonames.org/export/dump/readme.txt','sources':sources,'limitations':['Place information only; supplier city/country are from the catalog, not verified company locations.','cities15000 is not exhaustive. Multiple matching places remain ambiguous; never choose by population or supplier name.','Alternate names may include historical forms and transliterations; they retrieve candidates, not operational exposure.'],'locations':out}
args.output.parent.mkdir(parents=True,exist_ok=True)
args.output.write_text(json.dumps(data,indent=2,ensure_ascii=False)+'\n')
