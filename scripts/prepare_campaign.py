"""Read only the four explicitly authorised originals, never an answer key.
Writes a PRIVATE request payload outside Git. POST prepares, /start spends.
"""
import hashlib,json,re,sys
from pathlib import Path
from openpyxl import load_workbook
from app.importer import table,NEWS_MAP,SUPPLIER_MAP,SHEETS,iso,digest
from app.settings import ROOT

manifest=json.loads((ROOT/'evidence/input_manifest.json').read_text())
inputs=[];suppliers=[];materials=[];files=[]
for f in manifest['files']:
    if f['path'] not in ('inputs/supplier_training.xlsx','inputs/supplier_validation.xlsx','inputs/supplier_test.xlsx','inputs/supplier_liveday.xlsx'):continue
    data=(ROOT/f['path']).read_bytes();assert hashlib.sha256(data).hexdigest()==f['sha256']
    w=load_workbook(ROOT/f['path'],read_only=True,data_only=True)
    sheet=next(n for n in w.sheetnames if n in SHEETS);split=SHEETS[sheet];rows=table(w[sheet]);expected={'training':30,'validation':10,'test':20,'liveday':14}[split];assert len(rows)==expected
    for i,row in enumerate(rows):
        data={v:iso(row[k]) if k=='Published_Date' else str(row.get(k) or '') for k,v in NEWS_MAP.items()};data['context_id']='exercise'
        ref={k:row.get(k) for k in ('Is_Risk','Urgent','Affected_Supplier_IDs','Event_Type','Explanation')} if split in ('training','validation') else None
        inputs.append({'split':split,'row':i,'data':data,'reference':ref,'text_hash':digest(data),'reference_hash':digest(ref) if ref else None})
    if split=='training':
        for row in table(w['Suppliers']):
            s={v:row[k] for k,v in SUPPLIER_MAP.items()}
            if s['category'] in ('Logistics','Services'):s['inventory_days']=None
            suppliers.append(s)
        materials=table(w['Materials'])
    files.append({'path':f['path'],'filename':f['source_name'],'sha256':f['sha256'],'sheet':sheet,'split':split,'rows':len(rows),'labels':split in ('training','validation')});w.close()
assert len(inputs)==74 and len(suppliers)==40
payload={'name':'Comparativa nueva 74 × 2 · 19-10-2026','inputs':inputs,'suppliers':sorted(suppliers,key=lambda s:s['supplier_id']),'materials':materials,'manifest':{'files':files,'original_manifest_hash':digest(manifest),'catalog_hash':digest(sorted(suppliers,key=lambda s:s['supplier_id'])),'references_hash':digest([x['reference_hash'] for x in inputs]),'texts_hash':digest([x['text_hash'] for x in inputs]),'authorization':'Explicit user campaign: four originals, 74 inputs/profile, no answer key','independence':'training and validation already inspected; test/LiveDay unlabelled'}}
Path(sys.argv[1] if len(sys.argv)>1 else '.runtime/campaign-private-input.json').write_text(json.dumps(payload,ensure_ascii=False,default=str))
print(json.dumps({'inputs':len(inputs),'suppliers':len(suppliers),'splits':{s:sum(x['split']==s for x in inputs) for s in SHEETS.values()},'catalog_hash':payload['manifest']['catalog_hash'],'output':'private runtime file'}))
