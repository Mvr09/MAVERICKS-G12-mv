"""Authorized final phases. Private artifacts, frozen questions, persistent HTTP budget.

No answer key, prompt tuning, raw labels, credentials or protected answers in stdout.
"""

import argparse
from collections import Counter
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import time
import httpx
from openpyxl import load_workbook
from app.settings import settings, ROOT

BASE = "https://mavericks-g12-mv-production.up.railway.app"
BUDGET = "demo-final-200-20261017"
PHASES = [
    ("validation", "supplier_validation.xlsx", "News_Validation", 10),
    ("test", "supplier_test.xlsx", "News_Test", 20),
    ("liveday", "supplier_liveday.xlsx", "News_Feed", 14),
]
parser = argparse.ArgumentParser(
    description="Evaluación y ensayo final autorizados: máximo 200 peticiones nuevas compartidas"
)
parser.add_argument(
    "--start",
    action="store_true",
    help="Encolar fases autorizadas; sin este flag solo comprueba estado",
)
parser.add_argument(
    "--rss-after",
    action="store_true",
    help="Procesar RSS pendiente al terminar las tres fases dentro del mismo cupo",
)
args = parser.parse_args()
private = ROOT / "reports/private/final-rehearsal.json"
private.parent.mkdir(parents=True, exist_ok=True)
report = (
    json.loads(private.read_text())
    if private.exists()
    else {"phases": {}, "budget_id": BUDGET}
)
report["checked_at"] = datetime.now(timezone.utc).isoformat()
report.pop("check_error", None)


def save():
    private.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n")


def checked(response, expected=200):
    if response.status_code != expected:
        raise RuntimeError(f"HTTP_{response.status_code}")
    return response


def emit(value):
    print(json.dumps(value, ensure_ascii=False), flush=True)


def budget(client):
    status = checked(client.get("/api/admin/budgets")).json()
    return next(b for b in status["budgets"] if b["id"] == BUDGET)


def export_and_verify(content, expected_ids, split, supplier_ids):
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        rows = list(workbook.active.values)
        headers = list(rows[0])
        required = ["News_ID", "Is_Risk", "Urgent", "Affected_Supplier_IDs"]
        if not set(required) <= set(headers):
            raise RuntimeError("EXPORT_COLUMNS")
        indexes = {k: headers.index(k) for k in required}
        if [str(r[indexes["News_ID"]]) for r in rows[1:]] != expected_ids:
            raise RuntimeError("EXPORT_IDS_OR_ORDER")
        for row in rows[1:]:
            if any(row[indexes[k]] not in ("Y", "N") for k in ("Is_Risk", "Urgent")):
                raise RuntimeError("EXPORT_FLAGS")
            text = row[indexes["Affected_Supplier_IDs"]] or ""
            ids = text.split("; ") if text else []
            if ids != sorted(set(ids)) or any(s not in supplier_ids for s in ids):
                raise RuntimeError("EXPORT_SUPPLIER_IDS")
        if split == "test" and not {
            "Headline",
            "Summary",
            "Source",
            "Published_Date",
        } <= set(headers):
            raise RuntimeError("EXPORT_ORIGINAL_COLUMNS")
        return {
            "rows": len(rows) - 1,
            "ids_unique": len(expected_ids) == len(set(expected_ids)),
            "order_preserved": True,
            "required_columns": True,
            "original_columns_preserved": split == "test",
            "rss_excluded": True,
        }
    finally:
        workbook.close()


try:
    freeze = json.loads((ROOT / "reports/prompt-freeze.json").read_text())
    for name, digest in freeze["files"].items():
        if sha256((ROOT / name).read_bytes()).hexdigest() != digest:
            raise RuntimeError("FROZEN_CONFIGURATION_CHANGED")
    report["frozen_question_version"] = freeze["question_version"]
    report["frozen_policy_version"] = freeze["policy_version"]
    manifest = json.loads((ROOT / "evidence/input_manifest.json").read_text())
    allowed = {
        Path(entry["path"]).name: ROOT / entry["path"] for entry in manifest["files"]
    }
    with httpx.Client(
        base_url=BASE, headers={"Origin": BASE}, timeout=25, follow_redirects=False
    ) as client:
        config = checked(client.get("/api/config")).json()
        if (
            config["model"] != freeze["model"]
            or config["policy_version"] != freeze["policy_version"]
            or config.get("question_version") != freeze["question_version"]
        ):
            raise RuntimeError("DEPLOYED_CONFIGURATION_MISMATCH")
        if config["rss_paid_analysis"]:
            raise RuntimeError("PERIODIC_PAID_ANALYSIS_ENABLED")
        worker = checked(client.get("/api/worker")).json()
        if not worker["online"]:
            raise RuntimeError("WORKER_OFFLINE")
        if worker.get("compatible") is False:
            raise RuntimeError("WORKER_VERSION_MISMATCH")
        supplier_ids = {
            s["supplier_id"]
            for s in checked(client.get("/api/suppliers")).json()["items"]
        }
        if len(supplier_ids) != 40:
            raise RuntimeError("CATALOG_REQUIRES_40_SUPPLIERS")
        response = checked(
            client.post(
                "/api/session",
                json={"password": settings.admin_password.get_secret_value()},
            )
        )
        client.headers["X-CSRF-Token"] = response.json()["csrf"]
        try:
            report.setdefault(
                "execution_revision",
                report.get("revision")
                or checked(client.get("/healthz")).json()["revision"],
            )
            report["budget_before"] = report.get("budget_before") or budget(client)
            if report["budget_before"]["request_limit"] != 200:
                raise RuntimeError("BUDGET_LIMIT_MISMATCH")
            for split, filename, sheet_name, count in PHASES:
                phase = report["phases"].setdefault(split, {})
                if phase.get("completed"):
                    emit({"phase": split, "reused_existing_execution": True})
                    continue
                if not args.start:
                    continue
                if budget(client)["remaining_requests"] <= 0:
                    raise RuntimeError("REQUEST_BUDGET_EXHAUSTED")
                file = allowed[filename]
                content = file.read_bytes()
                expected_hash = next(
                    e["sha256"]
                    for e in manifest["files"]
                    if Path(e["path"]).name == filename
                )
                if sha256(content).hexdigest() != expected_hash:
                    raise RuntimeError("INPUT_MANIFEST_HASH_MISMATCH")
                workbook = load_workbook(
                    BytesIO(content), read_only=True, data_only=True
                )
                try:
                    rows = workbook[sheet_name].iter_rows(values_only=True)
                    headers = list(next(rows))
                    index = headers.index("News_ID")
                    expected_ids = [str(r[index]) for r in rows if r[index] is not None]
                finally:
                    workbook.close()
                if len(expected_ids) != count:
                    raise RuntimeError("INPUT_ROW_COUNT")
                imported = checked(
                    client.post(
                        "/api/imports",
                        files={
                            "file": (
                                filename,
                                content,
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            )
                        },
                    )
                ).json()
                bid = imported["batch_id"]
                phase["batch_id"] = bid
                before = next(
                    b
                    for b in checked(client.get("/api/batches")).json()
                    if b["id"] == bid
                )
                if not phase.get("execution_started_at") and before["done"]:
                    raise RuntimeError("FRESH_FINAL_MEASUREMENT_REQUIRED")
                phase.setdefault(
                    "execution_started_at", datetime.now(timezone.utc).isoformat()
                )
                save()
                if not before["started_at"]:
                    checked(client.post(f"/api/batches/{bid}/analyze"), 202)
                deadline = time.monotonic() + 330
                while True:
                    status = next(
                        b
                        for b in checked(client.get("/api/batches")).json()
                        if b["id"] == bid
                    )
                    phase["status"] = status
                    save()
                    if status["errors"]:
                        raise RuntimeError("ANALYSIS_ERROR_REQUIRES_DIAGNOSIS")
                    if status["done"] == count:
                        exported = checked(
                            client.get(f"/api/batches/{bid}/export.xlsx")
                        )
                        phase["export_received_at"] = datetime.now(
                            timezone.utc
                        ).isoformat()
                        phase["elapsed_to_export_seconds"] = round(
                            (
                                datetime.fromisoformat(phase["export_received_at"])
                                - datetime.fromisoformat(phase["execution_started_at"])
                            ).total_seconds(),
                            3,
                        )
                        phase["export_checks"] = export_and_verify(
                            exported.content, expected_ids, split, supplier_ids
                        )
                        output = ROOT / ".runtime/final" / f"{split}-resultados.xlsx"
                        output.parent.mkdir(parents=True, exist_ok=True)
                        output.write_bytes(exported.content)
                        phase["private_export"] = str(output.relative_to(ROOT))
                        phase["execution"] = checked(
                            client.get(f"/api/batches/{bid}/execution")
                        ).json()
                        phase["completed"] = True
                        phase["within_300_seconds"] = (
                            phase["elapsed_to_export_seconds"] <= 300
                            if split == "liveday"
                            else None
                        )
                        if split == "validation":
                            phase["metrics"] = checked(
                                client.post(f"/api/evaluations/{bid}")
                            ).json()
                        save()
                        emit(
                            {
                                "phase": split,
                                "rows": count,
                                "elapsed_to_export_seconds": phase[
                                    "elapsed_to_export_seconds"
                                ],
                                "network_attempts": len(phase["execution"]["calls"]),
                                "completed": True,
                            }
                        )
                        break
                    if time.monotonic() > deadline:
                        raise RuntimeError("FINAL_PHASE_TIMEOUT")
                    emit(
                        {
                            "phase": split,
                            "done": status["done"],
                            "total": count,
                            "errors": status["errors"],
                        }
                    )
                    time.sleep(1)
            if (
                args.start
                and args.rss_after
                and all(
                    report["phases"].get(s, {}).get("completed")
                    for s, _, _, _ in PHASES
                )
            ):
                report.setdefault("rss_jobs", [])
                for _ in range(9):
                    pending = [
                        a
                        for a in checked(client.get("/api/articles?context=rss")).json()
                        if a["state"] == "pending"
                    ]
                    available = budget(client)["remaining_requests"]
                    limit = min(10, len(pending), available // 6)
                    if limit < 1:
                        break
                    submitted = client.post("/api/analysis/rss", json={"limit": limit})
                    if submitted.status_code == 409:
                        break
                    jobs = checked(submitted, 202).json()["job_ids"]
                    report["rss_jobs"].extend(jobs)
                    save()
                    deadline = time.monotonic() + 120
                    while True:
                        states = [
                            checked(client.get("/api/jobs/" + jid)).json()["state"]
                            for jid in jobs
                        ]
                        if all(s in ("done", "error") for s in states):
                            break
                        if time.monotonic() > deadline:
                            raise RuntimeError("RSS_MONITOR_TIMEOUT")
                        time.sleep(2)
                    if "error" in states:
                        raise RuntimeError("RSS_ERROR_REQUIRES_DIAGNOSIS")
                report["rss_article_states"] = dict(
                    Counter(
                        a["state"]
                        for a in checked(client.get("/api/articles?context=rss")).json()
                    )
                )
            report["budget_after"] = budget(client)
            if report["budget_after"]["requests_used"] > 200:
                raise RuntimeError("BUDGET_INVARIANT_VIOLATION")
            report["revision"] = checked(client.get("/healthz")).json()["revision"]
            report["verified_revision"] = report["revision"]
            save()
            emit(
                {
                    "budget": report["budget_after"],
                    "completed_phases": [
                        s for s, p in report["phases"].items() if p.get("completed")
                    ],
                    "private_report": str(private.relative_to(ROOT)),
                }
            )
        finally:
            client.delete("/api/session")
except Exception as exc:
    report["check_error"] = (
        str(exc) if isinstance(exc, RuntimeError) else type(exc).__name__
    )
    save()
    emit({"check_error": report["check_error"]})
    raise SystemExit(1)
