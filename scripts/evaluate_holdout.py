"""Measures the deployed default profile on the independent holdout set.

Uses only the app's HTTP API (no app code, no database):
  login -> import inputs/supplier_holdout.xlsx -> analyse with the default
  profile -> wait -> download the export -> compare with
  inputs/supplier_holdout_labels.json -> delete the batch with the
  presenter's demo cleanup (/api/demo/preview + /api/demo/cleanup) ->
  check the public Ejercicio figures are back to their previous values.

  RADAR_URL=https://... RADAR_PASSWORD=... python scripts/evaluate_holdout.py
  python scripts/evaluate_holdout.py --preflight        # checks only, no import
  python scripts/evaluate_holdout.py --cleanup-batch ID # remove a leftover batch

Writes reports/holdout_eval_<date>.json and .md.

Why the cleanup is safe on the deployed version (schema 0010): the batch, its
articles, analyses, jobs, import rows and the incidents/alerts that only it
created are deleted. Request/consumption counters (ApiCall, budgets, grants)
are kept by design. Two cases would touch teammates' data, so they are
checked: an article of this batch joining an existing incident (cleanup then
re-creates that incident's alerts and drops its drafts) and schema 0011
(action plans reference incidents and are not deleted by the cleanup).
"""
import argparse, hashlib, io, json, math, os, re, sys, time
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "inputs/supplier_holdout.xlsx"
LABELS = ROOT / "inputs/supplier_holdout_labels.json"
EXPECTED_SCHEMA = "0010"  # cleanup verified for this schema only
RATE_PER_MILLION = 0.03941248835  # config/final_profile.json historical_rate
STOP = {"the", "and", "for", "with", "from", "that", "this", "after", "supplier",
        "supply", "news", "plant", "helixa"}  # app/incidents.py


def log(*a):
    print(*a, flush=True)


class Radar:
    def __init__(self, url, password):
        self.c = httpx.Client(base_url=url.rstrip("/"), timeout=60, follow_redirects=False)
        self.password, self.csrf = password, None

    def get(self, path, **params):
        r = self.c.get(path, params=params)
        r.raise_for_status()
        return r

    def send(self, method, path, **kw):
        r = self.c.request(method, path, headers={"X-CSRF-Token": self.csrf or ""}, **kw)
        if r.status_code >= 400:
            raise RuntimeError(f"{method} {path} -> {r.status_code}: {r.text[:300]}")
        return r

    def login(self):
        r = self.c.post("/api/session", json={"password": self.password})
        if r.status_code != 200:
            raise SystemExit(f"Login failed: HTTP {r.status_code} {r.text[:200]}")
        self.csrf = r.json()["csrf"]

    def logout(self):
        if self.csrf:
            try:
                self.send("DELETE", "/api/session")
            except Exception:
                pass


# ---------- public state snapshot (what teammates see) ----------
def digest(v):
    return hashlib.sha256(json.dumps(v, sort_keys=True, default=str).encode()).hexdigest()[:16]


def snapshot(api):
    ov = api.get("/api/overview", context="exercise").json()
    keep = ("received", "processed", "pending", "errors", "review", "review_model",
            "review_scope_resolved", "human_review_completed", "linked", "risk", "urgent",
            "supplier_count", "urgent_suppliers", "active_incidents", "review_rate",
            "automatic_rate", "latencies")
    alerts = api.get("/api/alerts", context="exercise").json()
    sup = api.get("/api/suppliers", context="exercise").json()["items"]
    incidents = sorted((i["id"], s["supplier_id"], i["state"], i["revision"], i["priority"])
                       for s in sup for i in s["incidents"])
    final = api.get("/api/model/final").json()
    metrics = api.get("/api/model/metrics", context="exercise").json()
    metrics.pop("updated_at", None)
    batches = sorted((b["id"], b["count"], b["done"], json.dumps(b["analysis_profiles"], sort_keys=True))
                     for b in api.get("/api/batches").json())
    return {
        "overview_exercise": {k: ov.get(k) for k in keep},
        "alerts_exercise": {"count": len(alerts["items"]), "unread": alerts["unread"],
                            "digest": digest(sorted((a["id"], a["acknowledged"], a["is_current"]) for a in alerts["items"]))},
        "suppliers_exercise": {"incidents": len(incidents), "digest": digest(incidents),
                               "status_digest": digest(sorted((s["supplier_id"], s["status"], s["priority"], s["urgent"], s["needs_review"]) for s in sup))},
        "batches": {"count": len(batches), "digest": digest(batches)},
        "model_final_digest": digest(final),
        "model_metrics_exercise_digest": digest(metrics),
        "campaigns_digest": digest(api.get("/api/campaigns").json()),
        "_incident_histories": [(s["supplier_id"], s["name"], i["id"], i["history"]) for s in sup for i in s["incidents"]],
    }


def compare(before, after):
    diffs = {}
    for k, v in before.items():
        if k.startswith("_"):
            continue
        if v != after.get(k):
            diffs[k] = {"before": v, "after": after.get(k)}
    return diffs


# ---------- incident-merge pre-flight (mirrors app/incidents.py) ----------
def tokens(text, supplier_name):
    excluded = set(re.findall(r"\w+", supplier_name.lower())) | STOP
    return {x for x in re.findall(r"[\w]+", text.lower()) if len(x) > 3 and x not in excluded}


def merge_risk(rows, histories):
    """Highest token overlap between a holdout article and an existing incident
    entry within 7 days. The app merges at >= 0.5 (same event type)."""
    worst = 0.0
    for _, name, _, history in histories:
        for h in history:
            if not h.get("tokens") or not h.get("published_date"):
                continue
            hd = datetime.fromisoformat(h["published_date"].replace("Z", "+00:00")).date()
            for r in rows:
                if abs((r["date"] - hd).days) > 7:
                    continue
                mine = tokens(r["headline"] + " " + r["summary"], name)
                theirs = set(h["tokens"])
                worst = max(worst, len(mine & theirs) / max(1, len(mine | theirs)))
    return round(worst, 3)


# ---------- metrics ----------
def ids(value):
    return {s.strip() for s in str(value or "").replace(",", ";").split(";") if s.strip()}


def binary(pairs):
    tp = sum(p == "Y" and e == "Y" for e, p in pairs)
    fp = sum(p == "Y" and e == "N" for e, p in pairs)
    fn = sum(p != "Y" and e == "Y" for e, p in pairs)
    tn = sum(p != "Y" and e == "N" for e, p in pairs)
    prec = tp / (tp + fp) if tp + fp else None
    rec = tp / (tp + fn) if tp + fn else None
    f1 = 2 * prec * rec / (prec + rec) if prec is not None and rec is not None and prec + rec else None
    return {"n": len(pairs), "tp": tp, "fp": fp, "fn": fn, "tn": tn,
            "precision": prec, "recall": rec, "f1": f1, "accuracy": (tp + tn) / len(pairs) if pairs else None}


def evaluate(labels, preds, view):
    rows = []
    for lab in labels:
        nid = lab["News_ID"]
        p = preds.get(nid)
        exp_ids, got_ids = ids(lab["Affected_Supplier_IDs"]), ids(p and p["Affected_Supplier_IDs"])
        v = view.get(nid, {})
        rows.append({
            "news_id": nid, "case_type": lab["case_type"],
            "expected": {"is_risk": lab["Is_Risk"], "urgent": lab["Urgent"], "suppliers": sorted(exp_ids)},
            "predicted": {"is_risk": p and p["Is_Risk"], "urgent": p and p["Urgent"], "suppliers": sorted(got_ids)} if p else None,
            "risk_ok": bool(p) and p["Is_Risk"] == lab["Is_Risk"],
            "urgent_ok": bool(p) and p["Urgent"] == lab["Urgent"],
            "suppliers_exact": bool(p) and got_ids == exp_ids,
            "row_exact": bool(p) and p["Is_Risk"] == lab["Is_Risk"] and p["Urgent"] == lab["Urgent"] and got_ids == exp_ids,
            "supplier_fp": sorted(got_ids - exp_ids), "supplier_fn": sorted(exp_ids - got_ids),
            "state_shown": v.get("state"), "review_scope_resolved": v.get("review_scope"),
            "review_reasons": v.get("review_reasons", []),
        })
    ok = [r for r in rows if r["predicted"]]
    sup_tp = sum(len(set(r["expected"]["suppliers"]) & set(r["predicted"]["suppliers"])) for r in ok)
    sup_fp = sum(len(r["supplier_fp"]) for r in ok)
    sup_fn = sum(len(r["supplier_fn"]) for r in rows)
    prec = sup_tp / (sup_tp + sup_fp) if sup_tp + sup_fp else None
    rec = sup_tp / (sup_tp + sup_fn) if sup_tp + sup_fn else None

    def block(sel):
        return {"n": len(sel),
                "risk_correct": sum(r["risk_ok"] for r in sel),
                "urgent_correct": sum(r["urgent_ok"] for r in sel),
                "suppliers_exact": sum(r["suppliers_exact"] for r in sel),
                "rows_exact": sum(r["row_exact"] for r in sel),
                "errors": [r["news_id"] for r in sel if not r["row_exact"]]}

    by_case = {c: block([r for r in rows if r["case_type"] == c])
               for c in sorted({r["case_type"] for r in rows})}
    auto = [r for r in rows if r["state_shown"] == "done" and not r["review_scope_resolved"]]
    scope = [r for r in rows if r["state_shown"] == "done" and r["review_scope_resolved"]]
    review = [r for r in rows if r["state_shown"] == "review"]
    return rows, {
        "n": len(rows), "predicted": len(ok),
        "risk": binary([(r["expected"]["is_risk"], r["predicted"]["is_risk"] if r["predicted"] else None) for r in rows]),
        "urgent": binary([(r["expected"]["urgent"], r["predicted"]["urgent"] if r["predicted"] else None) for r in rows]),
        "suppliers": {"pairs_expected": sum(len(r["expected"]["suppliers"]) for r in rows),
                      "tp": sup_tp, "fp": sup_fp, "fn": sup_fn, "precision": prec, "recall": rec,
                      "f1": 2 * prec * rec / (prec + rec) if prec and rec else None,
                      "exact_sets": sum(r["suppliers_exact"] for r in rows)},
        "rows_exact": sum(r["row_exact"] for r in rows),
        "by_case_type": by_case,
        "automatic_vs_review": {
            "automatic": block(auto),
            "review_scope_resolved": block(scope),
            "review": block(review),
            "note": "automatic = shown as done with no model doubt; review_scope_resolved = model asked for review but the app shows it resolved (no risk and no supplier involved); review = pending human review.",
        },
    }


def pct(x):
    return "n/d" if x is None else f"{x:.2f}"


def write_markdown(path, rep):
    m, meta = rep["metrics"], rep["run"]
    r, u, s = m["risk"], m["urgent"], m["suppliers"]
    av = m["automatic_vs_review"]
    lines = [
        f"# Holdout independiente · evaluación del perfil desplegado ({meta['date']})",
        "",
        "> **Muestra pequeña (30 noticias ficticias generadas por IA).** Un error mueve la precisión/recall varios puntos; "
        "las cifras orientan, no certifican. Ver construcción y límites en `docs/HOLDOUT.md`.",
        "",
        f"- Web: {meta['url']} · revisión `{meta['revision'][:7]}` · esquema {meta['schema']}",
        f"- Perfil por defecto: `{meta['default_profile']}` · preguntas `{meta['question_version']}` · JEV `{meta.get('jev')}`",
        f"- Lote importado `{meta['batch_id']}` · {meta['analysed']}/{m['n']} analizadas · errores {meta['article_errors']}",
        f"- Tiempo de lote (servidor): {pct(meta['batch_elapsed_seconds'])} s · cliente, importación→fin: {meta.get('client_seconds_import_to_done')} s · peticiones {meta['requests']} (HTTP 200: {meta['http_200']}) · "
        f"tokens entrada {meta['tokens']['input_tokens']} / salida {meta['tokens']['output_tokens']} · "
        f"coste aprox. {meta['approx_cost_usd']:.4f} $ (tasa histórica {RATE_PER_MILLION} $/M, no tarifa oficial)",
        f"- Limpieza: {rep['cleanup']['status']} · cifras públicas de Ejercicio tras limpiar: "
        + ("idénticas a las previas" if not rep["cleanup"].get("diffs") else "con diferencias (ver JSON)"),
        "",
        "## Resultados",
        "",
        "| Pregunta | TP | FP | FN | TN | Precision | Recall | F1 |",
        "|---|---|---|---|---|---|---|---|",
        f"| Is_Risk (n={r['n']}) | {r['tp']} | {r['fp']} | {r['fn']} | {r['tn']} | {pct(r['precision'])} | {pct(r['recall'])} | {pct(r['f1'])} |",
        f"| Urgent (n={u['n']}) | {u['tp']} | {u['fp']} | {u['fn']} | {u['tn']} | {pct(u['precision'])} | {pct(u['recall'])} | {pct(u['f1'])} |",
        f"| Proveedores (pares, {s['pairs_expected']} esperados) | {s['tp']} | {s['fp']} | {s['fn']} | – | {pct(s['precision'])} | {pct(s['recall'])} | {pct(s['f1'])} |",
        "",
        f"Conjuntos de proveedores exactos: **{s['exact_sets']}/{m['n']}** · filas exactas (las tres respuestas): **{m['rows_exact']}/{m['n']}**.",
        "",
        "## Por tipo de caso",
        "",
        "| case_type | n | Riesgo ok | Urgencia ok | Proveedores exactos | Filas exactas | Errores |",
        "|---|---|---|---|---|---|---|",
    ]
    for c, b in m["by_case_type"].items():
        lines.append(f"| {c} | {b['n']} | {b['risk_correct']}/{b['n']} | {b['urgent_correct']}/{b['n']} | "
                     f"{b['suppliers_exact']}/{b['n']} | {b['rows_exact']}/{b['n']} | {', '.join(b['errors']) or '–'} |")
    lines += ["", "## Automáticas vs revisión", "",
              "| Estado mostrado | n | Filas exactas | Errores |", "|---|---|---|---|"]
    for k, label in (("automatic", "Automática"), ("review_scope_resolved", "Duda del modelo resuelta por alcance"), ("review", "Revisión humana")):
        b = av[k]
        lines.append(f"| {label} | {b['n']} | {b['rows_exact']}/{b['n']} | {', '.join(b['errors']) or '–'} |")
    lines += ["", "## Errores fila a fila", "",
              "| News_ID | case_type | Esperado (riesgo/urgencia/proveedores) | Predicho | Estado |",
              "|---|---|---|---|---|"]
    for row in rep["rows"]:
        if row["row_exact"]:
            continue
        e, p = row["expected"], row["predicted"] or {}
        lines.append(f"| {row['news_id']} | {row['case_type']} | {e['is_risk']}/{e['urgent']}/{';'.join(e['suppliers']) or '–'} | "
                     f"{p.get('is_risk')}/{p.get('urgent')}/{';'.join(p.get('suppliers') or []) or '–'} | {row['state_shown']} |")
    lines += ["", "## Límites", "",
              "- Noticias y etiquetas generadas por IA a partir del README y el catálogo; sin revisión humana independiente.",
              "- 30 casos inequívocos por diseño: no mide casos dudosos ni ruido real de RSS; la distribución no es la del mundo real.",
              "- Una sola ejecución (sin repetir para medir variabilidad del modelo).",
              "- Las métricas se calculan sobre la exportación oficial del lote (columnas Is_Risk, Urgent, Affected_Supplier_IDs).", ""]
    path.write_text("\n".join(lines), encoding="utf-8")


# ---------- cleanup ----------
def cleanup(api, bid, expected_articles=None):
    for attempt in range(1, 6):
        prev = api.send("POST", "/api/demo/preview", json={"selection": ["batch:" + bid], "mode": "delete"}).json()
        counts = prev["counts"]
        if prev["blocked"]:
            if "en curso" in prev["blocked"] and attempt < 5:
                time.sleep(10)
                continue
            return {"status": "blocked", "reason": prev["blocked"], "counts": counts}
        if counts["batches"] != 1 or (expected_articles and counts["articles"] != expected_articles) or counts["drafts"]:
            return {"status": "refused", "reason": "preview counts differ from this batch", "counts": counts}
        try:
            res = api.send("POST", "/api/demo/cleanup", json={"confirmation": "LIMPIAR DEMO", "token": prev["token"]}).json()
            return {"status": "deleted", "counts": res["deleted"]}
        except RuntimeError as exc:
            if "han cambiado" in str(exc) and attempt < 5:
                time.sleep(3)
                continue
            return {"status": "error", "reason": str(exc), "counts": counts}
    return {"status": "error", "reason": "retries exhausted"}


# ---------- main ----------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--preflight", action="store_true", help="checks only, no import")
    ap.add_argument("--cleanup-batch", help="delete a leftover holdout batch and exit")
    ap.add_argument("--timeout", type=int, default=900)
    args = ap.parse_args()
    url, password = os.environ.get("RADAR_URL"), os.environ.get("RADAR_PASSWORD")
    if not url or not password:
        raise SystemExit("Set RADAR_URL and RADAR_PASSWORD")
    api = Radar(url, password)

    health = api.get("/healthz").json()
    schema = api.c.get("/readyz").json().get("schema")
    worker = api.get("/api/worker").json()
    config = api.get("/api/config").json()
    log("revision", health["revision"][:7], "schema", schema, "worker", worker["online"], worker["compatible"],
        "default profile", config["default_profile"])
    api.login()
    try:
        if args.cleanup_batch:
            log(json.dumps(cleanup(api, args.cleanup_batch), ensure_ascii=False))
            return
        if schema != EXPECTED_SCHEMA:
            raise SystemExit(f"Schema {schema} != {EXPECTED_SCHEMA}: cleanup not verified for this version (action plans). Not importing.")
        if not (worker["online"] and worker["compatible"]):
            raise SystemExit("Worker offline or on another revision. Not importing.")
        batches = api.get("/api/batches").json()
        if any(b["filename"] == XLSX.name for b in batches):
            raise SystemExit("A holdout batch already exists; remove it with --cleanup-batch first.")
        if any(b["queued"] or b["running"] for b in batches):
            raise SystemExit("Another batch is being analysed; try later.")

        labels = json.loads(LABELS.read_text(encoding="utf-8"))["labels"]
        wb = load_workbook(XLSX, read_only=True)
        sheet = list(wb["News_Test"].iter_rows(values_only=True))
        wb.close()
        head = sheet[0]
        rows = [dict(zip(head, r)) for r in sheet[1:] if r[0]]
        assert all(r.get(k) in (None, "") for r in rows for k in ("Is_Risk", "Urgent", "Affected_Supplier_IDs")), "answers present in xlsx"
        assert [r["News_ID"] for r in rows] == [l["News_ID"] for l in labels]
        news = [{"date": r["Published_Date"].date(), "headline": r["Headline"], "summary": r["Summary"]} for r in rows]

        before = snapshot(api)
        risk = merge_risk(news, before["_incident_histories"])
        log("existing exercise incidents", before["suppliers_exercise"]["incidents"], "max merge overlap", risk)
        if risk >= 0.35:
            raise SystemExit(f"An article could join an existing incident (overlap {risk}); not importing.")
        budgets_before = {b["id"]: b["requests_used"] for b in api.get("/api/admin/budgets").json()["budgets"]}
        if args.preflight:
            log("preflight ok; nothing imported")
            return

        imp = api.send("POST", "/api/imports", files={"file": (XLSX.name, XLSX.read_bytes(),
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}).json()
        bid = imp["batch_id"]
        log("imported", imp)
        report = {"cleanup": {"status": "not run"}}
        try:
            if imp.get("idempotent") or imp.get("rows") != len(labels):
                raise RuntimeError(f"unexpected import result {imp}")
            started = time.time()
            jobs = api.send("POST", f"/api/batches/{bid}/analyze").json()["job_ids"]
            log("analysis queued", len(jobs), "jobs")
            while True:
                b = next(x for x in api.get("/api/batches").json() if x["id"] == bid)
                log(f"  {int(time.time() - started)}s done {b['done']} errors {b['errors']} queued {b['queued']} running {b['running']}")
                if not b["pending"] and not b["queued"] and not b["running"]:
                    break
                if time.time() - started > args.timeout:
                    raise RuntimeError("timeout waiting for the analysis")
                time.sleep(5)
            for _ in range(10):  # completed_at is stamped just after the last job
                if b["elapsed_seconds"] is not None:
                    break
                time.sleep(2)
                b = next(x for x in api.get("/api/batches").json() if x["id"] == bid)
            articles =api.get("/api/articles", origin="batch:" + bid).json()
            view = {a["news_id"]: {"state": a["state"], "review_scope": bool(a.get("review_scope")),
                                   "review_reasons": (a.get("analysis") or {}).get("review_reasons", []),
                                   "id": a["id"]} for a in articles}
            preds, export_note = {}, "official export"
            exp = api.c.get(f"/api/batches/{bid}/export.xlsx")
            if exp.status_code == 200:
                ws = load_workbook(io.BytesIO(exp.content), read_only=True)["Resultados"]
                er = list(ws.iter_rows(values_only=True))
                preds = {r[er[0].index("News_ID")]: dict(zip(er[0], r)) for r in er[1:]}
            else:
                export_note = f"export blocked (HTTP {exp.status_code}); predictions from article view"
                for a in articles:
                    an = a.get("analysis") or {}
                    if an.get("analysis_status") != "error" and an.get("is_risk"):
                        preds[a["news_id"]] = {"Is_Risk": an["is_risk"], "Urgent": an["urgent"],
                                               "Affected_Supplier_IDs": "; ".join(an.get("affected_supplier_ids", []))}
            calls = api.get(f"/api/batches/{bid}/execution").json()["calls"]
            tok = {k: sum(c["usage"].get(k, 0) for c in calls) for k in ("input_tokens", "output_tokens")}
            mine = {v["id"] for v in view.values()}
            shared = []
            for sid, _, iid, history in snapshot_histories(api):
                aids = {h.get("article_id") for h in history if h.get("article_id")}
                if aids & mine and aids - mine:
                    shared.append({"supplier_id": sid, "incident_id": iid})
            rows_eval, metrics = evaluate(labels, preds, view)
            prof = (b.get("analysis_profiles") or [{}])[0]
            report.update({
                "run": {"date": datetime.now(timezone.utc).date().isoformat(), "url": url, "revision": health["revision"],
                        "schema": schema, "default_profile": config["default_profile"],
                        "question_version": prof.get("questions"), "review_version": prof.get("review"),
                        "jev": (config.get("default_components") or {}).get("model"),
                        "profile_components": config.get("default_components"),
                        "batch_id": bid, "analysed": b["done"], "article_errors": b["errors"],
                        "mixed_profiles": b["mixed_profiles"], "batch_elapsed_seconds": b["elapsed_seconds"],
                        "client_seconds_import_to_done": round(time.time() - started, 1),
                        "requests": len(calls), "http_200": sum(c["http_status"] == 200 for c in calls),
                        "retries": sum((c["attempt"] or 1) > 1 for c in calls), "tokens": tok,
                        "approx_cost_usd": (tok["input_tokens"] + tok["output_tokens"]) / 1e6 * RATE_PER_MILLION,
                        "predictions_source": export_note},
                "metrics": metrics, "rows": rows_eval,
                "shared_incidents": shared,
                "warning": "Small sample (30 AI-generated fictional articles). One error moves precision/recall by several points.",
            })
        finally:
            report["cleanup"] = cleanup(api, bid, len(labels))
            log("cleanup", report["cleanup"])
        after = snapshot(api)
        report["cleanup"]["diffs"] = compare(before, after)
        budgets_after = {b["id"]: b["requests_used"] for b in api.get("/api/admin/budgets").json()["budgets"]}
        report["cleanup"]["kept_by_design"] = {
            "request_budget_delta": {k: budgets_after.get(k, 0) - v for k, v in budgets_before.items() if budgets_after.get(k, 0) != v},
            "note": "ApiCall rows, request budgets and grant counters keep this run's consumption (the app never deletes them)."}
        log("public figures changed after cleanup:", list(report["cleanup"]["diffs"]) or "none")
        if "metrics" not in report:
            raise SystemExit("Run failed before metrics; batch cleanup: " + report["cleanup"]["status"])
        day = report["run"]["date"]
        out = ROOT / "reports"
        (out / f"holdout_eval_{day}.json").write_text(json.dumps(report, indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8")
        write_markdown(out / f"holdout_eval_{day}.md", report)
        m = report["metrics"]
        log("risk", m["risk"], "\nurgent", m["urgent"], "\nsuppliers", m["suppliers"], "\nrows exact", m["rows_exact"])
    finally:
        api.logout()


def snapshot_histories(api):
    sup = api.get("/api/suppliers", context="exercise").json()["items"]
    return [(s["supplier_id"], s["name"], i["id"], i["history"]) for s in sup for i in s["incidents"]]


if __name__ == "__main__":
    main()
