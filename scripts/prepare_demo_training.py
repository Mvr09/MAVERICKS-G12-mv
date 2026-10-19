"""Prepare a private three-row training subset. Never opens held-out workbooks."""
from pathlib import Path
from openpyxl import load_workbook,Workbook
ROOT=Path(__file__).resolve().parents[1]
def main():
 output=ROOT/'.runtime/Demo_training_3_filas.xlsx'
 if output.exists():
  print('Subconjunto privado existente conservado; no sobrescrito.');return
 source=load_workbook(ROOT/'inputs/supplier_training.xlsx',read_only=True,data_only=True)
 try:
  sheet=source['News'];it=sheet.iter_rows(values_only=True);headers=list(next(it));columns=['News_ID','Published_Date','Source','Headline','Summary','Language'];assert all(c in headers for c in columns)
  selected={};wanted=['N001','N002','N005']
  for row in it:
   news_id=str(row[headers.index('News_ID')])
   if news_id in wanted:selected[news_id]=[row[headers.index(c)]for c in columns]
  assert set(selected)==set(wanted)
  book=Workbook();out=book.active;out.title='News';out.append(columns)
  for key in wanted:out.append(selected[key])
  output.parent.mkdir(exist_ok=True);book.save(output)
 finally:source.close()
 print('Preparadas3filas training sin etiquetas; originales intactos.')
if __name__=='__main__':main()
