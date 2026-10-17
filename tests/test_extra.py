from datetime import timedelta
from io import BytesIO
import json
import pytest, httpx
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from app.db import Session
from app.models import Incident, Job, Grant, now
from app.importer import import_excel, export_batch, iso
from app.analysis import process_article
from app.drafts import make_draft
from app.queue import claim
from tests.helpers import synthetic, answer


def test_corrupt_macro_and_missing_schema():
    with Session() as db:
        with pytest.raises(ValueError, match="corrupto"):
            import_excel(db, b"not-an-excel", "x.xlsx")
        with pytest.raises(ValueError, match="XLSX"):
            import_excel(db, b"anything", "macro.xlsm")
        w = Workbook()
        w.active.title = "News"
        w.active.append(["News_ID"])
        w.active.append(["missing-columns"])
        b = BytesIO()
        w.save(b)
        with pytest.raises(ValueError, match="Faltan columnas"):
            import_excel(db, b.getvalue(), "x.xlsx")


def test_dates_normalize_utc():
    assert iso("2026-10-17") == "2026-10-17T00:00:00+00:00"
    assert iso("2026-10-17T14:00:00+02:00") == "2026-10-17T12:00:00+00:00"


def test_low_confidence_exports_binary_and_requires_review():
    with Session() as db:
        a = synthetic(db)
        result = process_article(
            db,
            a.id,
            httpx.MockTransport(
                lambda r: httpx.Response(
                    200,
                    json=answer(
                        json.loads(r.content), linked=("S001",), confidence=0.4
                    ),
                )
            ),
        )
        assert result["analysis_status"] == "review"
        assert result["assessments"][0]["priority"] > 0
        inc = db.scalar(select(Incident))
        assert inc
        with pytest.raises(ValueError, match="Revisa"):
            make_draft(
                db,
                db.get(
                    __import__("app.models", fromlist=["Supplier"]).Supplier, "S001"
                ),
                inc,
                "supplier",
            )
        assert load_workbook(BytesIO(export_batch(db, a.batch_id))).active.max_row == 2


def test_drafts_preserve_unknown_stock_and_no_external_send():
    from app.models import Supplier

    with Session() as db:
        a = synthetic(db)
        process_article(
            db,
            a.id,
            httpx.MockTransport(
                lambda r: httpx.Response(
                    200, json=answer(json.loads(r.content), linked=("S001",))
                )
            ),
        )
        inc = db.scalar(select(Incident))
        s = db.get(Supplier, "S001")
        draft = make_draft(db, s, inc, "alternate")
        assert "no constituye un pedido" in draft
        assert "No disponemos de confirmación de stock" in draft
        assert "Confirmación de continuidad" in make_draft(db, s, inc, "supplier")


def test_exhausted_lease_becomes_visible_error():
    with Session() as db:
        j = Job(
            key="mock-exhausted",
            kind="rss",
            payload={"source_id": "ema"},
            state="running",
            attempts=3,
            lease_until=now() - timedelta(seconds=1),
        )
        db.add(j)
        db.commit()
        assert claim(db) is None
        db.commit()
        db.refresh(j)
        assert j.state == "error"
        assert j.error == "LEASE_ATTEMPTS_EXHAUSTED"


def test_diagnosed_retry_keeps_attempt_budget():
    from app.jev import JevClient, ProviderError
    from reference.jev_requests import stage_a
    from app.models import Supplier, ApiCall
    import time

    with Session() as db:
        a = synthetic(db)
        payload = stage_a(a.data, [s.data for s in db.scalars(select(Supplier))])
        client = JevClient(db, a.id, httpx.MockTransport(lambda r: httpx.Response(401)))
        with pytest.raises(ProviderError, match="401"):
            client.ask(payload, "a", time.monotonic() + 60)
        # Blind retry does not issue another call.
        with pytest.raises(ProviderError, match="401"):
            client.ask(payload, "a", time.monotonic() + 60)
        assert len(db.scalars(select(ApiCall)).all()) == 1
        diagnosed = JevClient(
            db,
            a.id,
            httpx.MockTransport(
                lambda r: httpx.Response(200, json=answer(json.loads(r.content)))
            ),
            diagnostic_ack_at=now().isoformat(),
        )
        diagnosed.ask(payload, "a", time.monotonic() + 60)
        assert len(db.scalars(select(ApiCall)).all()) == 2


def test_expired_authorization_prevents_actual_http_request():
    from app.models import GrantArticle, ApiCall, Supplier
    from app.jev import JevClient, ProviderError
    from reference.jev_requests import stage_a
    import time

    with Session() as db:
        a = synthetic(db)
        grant = Grant(
            id="mock-expired",
            split="training",
            article_limit=1,
            used=1,
            expires_at=now() - timedelta(seconds=1),
        )
        db.add(grant)
        db.flush()
        db.add(GrantArticle(article_id=a.id, grant_id=grant.id))
        db.commit()
        payload = stage_a(a.data, [s.data for s in db.scalars(select(Supplier))])
        # No transport mock: the expired grant must stop before an HTTP client is opened.
        with pytest.raises(ProviderError, match="AUTHORIZATION_EXPIRED"):
            JevClient(db, a.id).ask(payload, "a", time.monotonic() + 60)
        assert db.scalars(select(ApiCall)).all() == []
