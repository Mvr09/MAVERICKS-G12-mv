import json
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from io import BytesIO
import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import select, func
from app.db import Session
from app.models import (
    ConsumptionBudget,
    Grant,
    GrantArticle,
    ApiCall,
    Article,
    Analysis,
    Incident,
    Supplier,
    ImportRow,
    Batch,
    EvaluationLabel,
    now,
)
from app.jev import reserve_request, JevClient, ProviderError, grant_requires_diagnosis
from app.analysis import process_article
from app.drafts import make_draft
from app.importer import export_batch
from app.evaluation import evaluate_batch
from reference.jev_requests import stage_a
from tests.helpers import synthetic, answer


def attach(db, a, budget_id="MOCK-SHARED", grant_id="MOCK-GRANT", limit=3):
    if not db.get(ConsumptionBudget, budget_id):
        db.add(
            ConsumptionBudget(
                id=budget_id, request_limit=limit, requests_used=0, window_seconds=2700
            )
        )
        db.flush()
    if not db.get(Grant, grant_id):
        db.add(
            Grant(
                id=grant_id,
                split="training",
                article_limit=10,
                used=1,
                budget_id=budget_id,
            )
        )
        db.flush()
    db.add(GrantArticle(article_id=a.id, grant_id=grant_id))
    db.commit()
    return db.get(Grant, grant_id)


def test_shared_request_budget_is_atomic_across_workers():
    with Session() as db:
        a = synthetic(db)
        attach(db, a)

    def attempt(_):
        with Session() as db:
            try:
                reserve_request(db, db.get(Grant, "MOCK-GRANT"))
                db.commit()
                return True
            except ProviderError:
                db.rollback()
                return False

    with ThreadPoolExecutor(max_workers=6) as pool:
        assert sum(pool.map(attempt, range(6))) == 3
    with Session() as db:
        b = db.get(ConsumptionBudget, "MOCK-SHARED")
        assert b.requests_used == 3 and b.started_at


def test_retries_consume_shared_budget_and_stop_before_next_http_call():
    requests = []

    def failed(r):
        requests.append(1)
        return httpx.Response(503)

    with Session() as db:
        a = synthetic(db)
        attach(db, a)
        payload = stage_a(a.data, [s.data for s in db.scalars(select(Supplier))])
        with pytest.raises(ProviderError):
            JevClient(db, a.id, httpx.MockTransport(failed), sleep=lambda _: None).ask(
                payload, "a", time.monotonic() + 60
            )
        assert len(requests) == 3
        b = synthetic(db, headline="MOCK second independent article")
        attach(db, b, grant_id="MOCK-SECOND-GRANT")
        payload = stage_a(b.data, [s.data for s in db.scalars(select(Supplier))])
        with pytest.raises(ProviderError, match="REQUEST_BUDGET_EXHAUSTED"):
            JevClient(db, b.id, httpx.MockTransport(failed)).ask(
                payload, "a", time.monotonic() + 60
            )
        assert len(requests) == 3
        assert (
            db.scalar(
                select(func.count())
                .select_from(ApiCall)
                .where(ApiCall.budget_id == "MOCK-SHARED")
            )
            == 3
        )


def test_shared_budget_window_and_authentication_circuit_survive_grant_change():
    with Session() as db:
        a = synthetic(db)
        g = attach(db, a)
        budget = db.get(ConsumptionBudget, g.budget_id)
        budget.started_at = now() - timedelta(seconds=2701)
        db.commit()
        with pytest.raises(ProviderError, match="BUDGET_WINDOW_EXPIRED"):
            reserve_request(db, g)
        db.rollback()
        budget.started_at = None
        db.add(
            ApiCall(
                article_id=a.id,
                stage="a",
                attempt=1,
                model="MOCK",
                status=401,
                duration=0,
                usage={},
                budget_id=g.budget_id,
            )
        )
        b = synthetic(db, headline="MOCK different grant")
        attach(db, b, grant_id="MOCK-SECOND-GRANT")
        assert grant_requires_diagnosis(db, "MOCK-SECOND-GRANT")


def test_draft_uses_validated_fragment_from_latest_published_evidence():
    with Session() as db:
        a = synthetic(db)
        process_article(
            db,
            a.id,
            httpx.MockTransport(
                lambda r: httpx.Response(
                    200, json=answer(json.loads(r.content), linked=("S001",))
                )
            ),
        )
        inc = db.scalar(select(Incident))
        evidence = inc.history[0]["explanation"]
        inc.history = [
            *inc.history,
            {
                "article_id": "MOCK-OLD",
                "published_date": "2025-01-01",
                "assessment": {"needs_review": True},
                "explanation": "OLD MUST NOT REPLACE CURRENT",
            },
        ]
        draft = make_draft(db, db.get(Supplier, "S001"), inc, "supplier")
        assert evidence in draft and "OLD MUST NOT REPLACE CURRENT" not in draft
        inc.history = [
            {"published_date": "2026-10-17", "assessment": {"needs_review": False}}
        ]
        with pytest.raises(ValueError, match="Falta evidencia"):
            make_draft(db, db.get(Supplier, "S001"), inc, "supplier")


def test_test_export_preserves_input_columns_and_excludes_evaluation_labels():
    with Session() as db:
        a = synthetic(db)
        batch = db.get(Batch, a.batch_id)
        batch.split = "test"
        batch.export_headers = ["News_ID", "Headline", "Published_Date"]
        db.add(
            ImportRow(
                article_id=a.id,
                data={
                    "News_ID": a.news_id,
                    "Headline": "=MOCK FORMULA STORED AS TEXT",
                    "Published_Date": "2026-10-17",
                },
            )
        )
        process_article(
            db,
            a.id,
            httpx.MockTransport(
                lambda r: httpx.Response(
                    200, json=answer(json.loads(r.content), linked=())
                )
            ),
        )
        sheet = load_workbook(BytesIO(export_batch(db, batch.id))).active
        headers = [c.value for c in sheet[1]]
        assert headers == [
            "News_ID",
            "Headline",
            "Published_Date",
            "Is_Risk",
            "Urgent",
            "Affected_Supplier_IDs",
        ]
        assert (
            sheet.cell(2, 2).value == "=MOCK FORMULA STORED AS TEXT"
            and sheet.cell(2, 2).data_type == "s"
        )
        assert not {"Explanation", "Event_Type"} & set(headers)


def test_evaluation_denominators_exact_set_and_missing_labels():
    with Session() as db:
        a = synthetic(db)
        db.add(
            EvaluationLabel(
                article_id=a.id,
                data={"Is_Risk": "Y", "Urgent": "Y", "Affected_Supplier_IDs": "S001"},
            )
        )
        db.commit()
        process_article(
            db,
            a.id,
            httpx.MockTransport(
                lambda r: httpx.Response(
                    200, json=answer(json.loads(r.content), linked=("S001",))
                )
            ),
        )
        result = evaluate_batch(db, a.batch_id)
        assert (
            result["n"] == 1 and result["risk"]["tp"] == 1 and result["risk"]["f1"] == 1
        )
        assert (
            result["suppliers"]["exact_set_accuracy"] == 1
            and result["suppliers"]["f1"] == 1
        )
        db.get(EvaluationLabel, a.id).data = {"Is_Risk": "Y"}
        db.commit()
        with pytest.raises(ValueError, match="etiquetas binarias"):
            evaluate_batch(db, a.batch_id)


def test_manual_recovery_disables_sound_and_execution_metadata_is_private():
    from fastapi.testclient import TestClient
    from app.main import app
    from app.models import Alert
    from tests.test_operations import login

    with Session() as db:
        a = synthetic(db)
        process_article(
            db,
            a.id,
            httpx.MockTransport(
                lambda r: httpx.Response(
                    200, json=answer(json.loads(r.content), linked=("S001",))
                )
            ),
        )
        inc = db.scalar(select(Incident))
        iid = inc.id
        bid = a.batch_id
    with TestClient(app) as c:
        assert c.get("/api/admin/budgets").status_code == 401
        assert c.get(f"/api/batches/{bid}/execution").status_code == 401
        headers = login(c)
        assert c.get(f"/api/batches/{bid}/execution").json()["calls"][0]["usage"] == {
            "input_tokens": 123,
            "output_tokens": 45,
        }
        response = c.post(
            f"/api/incidents/{iid}/resolve",
            headers=headers,
            json={"reason": "MOCK recuperación confirmada para ensayo"},
        )
        assert response.status_code == 200
    with Session() as db:
        alert = db.scalar(select(Alert).where(Alert.incident_id == iid))
        assert (
            alert.data["state"] == "resolved" and alert.data["sound_eligible"] is False
        )


def test_paid_jobs_require_workers_with_the_current_version():
    from fastapi.testclient import TestClient
    from app.main import app
    from app.models import WorkerPresence
    from tests.test_operations import login

    with Session() as db:
        db.add(
            WorkerPresence(
                id="MOCK-OLD-WORKER", heartbeat=now(), revision="OLD-UNBOUNDED-VERSION"
            )
        )
        db.commit()
    with TestClient(app) as c:
        headers = login(c)
        status = c.get("/api/worker").json()
        assert status["online"] is True and status["compatible"] is False
        assert (
            c.post("/api/analysis/rss", headers=headers, json={"limit": 1}).status_code
            == 409
        )
