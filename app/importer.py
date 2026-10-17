import hashlib, io, json, re, zipfile
from datetime import datetime, date, timezone
from openpyxl import load_workbook, Workbook
from sqlalchemy import select
from app.models import Supplier, Material, Article, Batch, EvaluationLabel
from app.settings import settings

LABELS = {'Is_Risk', 'Urgent', 'Affected_Supplier_IDs', 'Event_Type', 'Explanation'}
SHEETS = {'News':'training','News_Validation':'validation','News_Test':'test','News_Feed':'liveday'}
SUPPLIER_MAP = dict(zip(('Supplier_ID','Supplier_Name','Country','City','Category','Material_Supplied','Criticality','Single_Source','Approved_Alternate_ID','Lead_Time_Days','Inventory_Cover_Days','Share_of_Category_Spend_%','Annual_Spend_EUR','Quality_Score'), ('supplier_id','name','country','city','category','material','criticality','single_source','alternate_id','lead_days','inventory_days','share_category_pct','annual_spend','quality_score')))
NEWS_MAP = {'News_ID':'news_id','Published_Date':'published_date','Source':'source','Headline':'headline','Summary':'summary','Language':'language'}

def digest(value): return hashlib.sha256(json.dumps(value,sort_keys=True,ensure_ascii=False,default=str).encode()).hexdigest()
def iso(value):
    if value is None: raise ValueError('Falta fecha de publicación')
    if isinstance(value, (date, datetime)): return value.isoformat()
    try: return datetime.fromisoformat(str(value).replace('Z','+00:00')).isoformat()
    except ValueError: raise ValueError('Fecha de publicación inválida')

def table(ws):
    rows = ws.iter_rows(values_only=True)
    headers = list(next(rows))
    if len(set(headers)) != len(headers): raise ValueError(f'{ws.title}: columnas repetidas')
    return [dict(zip(headers,row)) for row in rows if any(v is not None for v in row)]

def import_excel(db, data: bytes, filename: str):
    if not filename.lower().endswith('.xlsx'): raise ValueError('Solo se admite XLSX sin macros')
    if len(data)>settings.max_upload_bytes: raise ValueError('Archivo superior a 10 MB')
    file_hash=hashlib.sha256(data).hexdigest()
    existing=db.scalar(select(Batch).where(Batch.file_hash==file_hash))
    if existing: return existing, {'idempotent':True,'rows':existing.count}
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(x.file_size for x in z.infolist())>80_000_000: raise ValueError('XLSX expandido demasiado grande')
            if any('vbaproject' in x.filename.lower() for x in z.infolist()): raise ValueError('Macros no permitidas')
        w=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    except (zipfile.BadZipFile,KeyError,IOError): raise ValueError('El archivo XLSX está corrupto')
    names=[n for n in w.sheetnames if n in SHEETS]
    if len(names)!=1: raise ValueError('Se requiere exactamente una hoja News, News_Validation, News_Test o News_Feed')
    name=names[0]; rows=table(w[name])
    if not rows or len(rows)>settings.max_import_rows: raise ValueError('Cantidad de filas inválida (máximo 5.000)')
    required={'News_ID','Published_Date','Source','Headline','Summary'}
    if not required.issubset(rows[0]): raise ValueError('Faltan columnas: '+', '.join(sorted(required-set(rows[0]))))
    ids=[str(r['News_ID']) for r in rows]
    if len(set(ids))!=len(ids) or any(x in ('None','') for x in ids): raise ValueError('News_ID ausente o repetido')
    if 'Suppliers' in w.sheetnames:
        suppliers=table(w['Suppliers'])
        for r in suppliers:
            if not set(SUPPLIER_MAP).issubset(r): raise ValueError('Esquema Suppliers incompleto')
            s={v:r[k] for k,v in SUPPLIER_MAP.items()}
            for field in ('lead_days','inventory_days','annual_spend','quality_score','share_category_pct'):
                v=s[field]
                if v is not None and (not isinstance(v,(int,float)) or v<0): raise ValueError(f'{s["supplier_id"]}: {field} inválido')
            if s['share_category_pct'] is not None and s['share_category_pct']>100: raise ValueError('Porcentaje de gasto fuera de rango')
            if s['criticality'] not in ('High','Medium','Low') or s['single_source'] not in ('Y','N'): raise ValueError('Criticidad o fuente única inválida')
            if s['category'] in ('Logistics','Services'): s['inventory_days']=None
            db.merge(Supplier(id=s['supplier_id'],data=s))
    if 'Materials' in w.sheetnames:
        for r in table(w['Materials']):
            required_m={'Material_ID','Material','Primary_Supplier_ID','Alternate_Supplier_ID','Used_In_Products','Monthly_Requirement','Unit','Category'}
            if not required_m.issubset(r): raise ValueError('Esquema Materials incompleto')
            m={'material_id':r['Material_ID'],'name':r['Material'],'category':r['Category'],'primary_id':r['Primary_Supplier_ID'],'alternate_id':r['Alternate_Supplier_ID'],'products':[x.strip() for x in re.split(r'[;\n]',str(r['Used_In_Products'] or '')) if x.strip()],'monthly_requirement':r['Monthly_Requirement'],'unit':r['Unit']}
            db.merge(Material(id=m['material_id'],data=m))
    normalized=[]
    for i,r in enumerate(rows):
        a={v:str(r.get(k) or '') for k,v in NEWS_MAP.items()}
        a['published_date']=iso(r['Published_Date']); a['context_id']='exercise'
        a['language']=a['language'] or 'en'
        if not a['headline'].strip(): raise ValueError(f'Fila {i+2}: falta titular')
        normalized.append(a)
    batch=Batch(file_hash=file_hash,split=SHEETS[name],count=len(rows),logical_date=max(a['published_date'] for a in normalized))
    db.add(batch); db.flush()
    for i,(r,a) in enumerate(zip(rows,normalized)):
        art=Article(origin_key=f'{batch.id}:{a["news_id"]}',batch_id=batch.id,context='exercise',row_order=i,news_id=a['news_id'],content_hash=digest(a),data=a)
        db.add(art); db.flush()
        labels={k:r[k] for k in LABELS if k in r and r[k] is not None}
        if labels: db.add(EvaluationLabel(article_id=art.id,data=labels))
    w.close(); db.flush()
    return batch, {'idempotent':False,'rows':len(rows),'sheet':name,'split':batch.split}

def export_batch(db,batch_id):
    from app.models import Analysis
    articles=db.scalars(select(Article).where(Article.batch_id==batch_id).order_by(Article.row_order)).all()
    if not articles: raise ValueError('Lote vacío o inexistente')
    w=Workbook(); s=w.active; s.title='Resultados'
    s.append(['News_ID','Is_Risk','Urgent','Affected_Supplier_IDs'])
    for a in articles:
        result=db.scalar(select(Analysis).where(Analysis.article_id==a.id).order_by(Analysis.created_at.desc()))
        if not result or result.result['analysis_status']=='error' or a.state not in ('done','review'): raise ValueError('Exportación bloqueada: quedan filas pendientes o con error')
        r=result.result
        values=[a.news_id,r['is_risk'],r['urgent'],'; '.join(sorted(set(r['affected_supplier_ids'])))]
        s.append(values)
        for c in s[s.max_row]: c.data_type='s'
    out=io.BytesIO();w.save(out);return out.getvalue()
